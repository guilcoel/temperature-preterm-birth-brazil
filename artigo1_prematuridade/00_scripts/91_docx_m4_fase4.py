# 91_docx_m4_fase4.py — Fase 4, marco M4: monta os .docx da submissão (A1-127). python-docx (sem pandoc nem officer).
# Entradas (do script 90 e do 87 v4; SHA-256 conferido contra os logs):
#   manuscrito/m4/manuscrito_<RODADA>.md, suplemento_<RODADA>.md, pagina_titulo_<RODADA>.md, highlights_<RODADA>.md,
#   legendas_<RODADA>.md; tabelas/*_v4_<RODADA>.csv; manuscrito/referencias_m4.csv (verificadas pelo subagente);
#   figuras/fig0N_*_v2_<RODADA>.pdf; da v7, a Figura S1 (PDF e PNG) e a Tabela S6 do script 93 v2 (SHA-256 contra o log do 93).
# Regras (mensagem do chat de 19/10, seção 4; errata E4):
#   - texto sem justificar, sem hifenização, coluna única, numeração de linhas contínua, Arial;
#   - citações [@chave] ou [@a; @b] viram números sobrescritos, na ordem da primeira citação; a chave sem referência verificada
#     vira "[ref]" e é listada no log;
#   - tabelas de três linhas (acima do cabeçalho, abaixo dele e no fim), sem bordas verticais nem sombreamento, conferidas no XML.
# Saídas (fora do git): manuscrito/m4/*.docx e a pasta Submissao/ (só os arquivos da revista). Log logs/docx_m4_fase4_<RODADA>.md.
# Uso (de artigo1_prematuridade/):  RODADA=AAAAMMDD python3 -I 00_scripts/91_docx_m4_fase4.py
import csv, datetime, hashlib, os, re, shutil, subprocess, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = "/Volumes/Forja/Doc/artigos_perinatais_clima"
ESTE = "artigo1_prematuridade/00_scripts/91_docx_m4_fase4.py"


def git(*a):
    return subprocess.run(["git", "-C", REPO, *a], capture_output=True, text=True, check=True).stdout.strip()


def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


rodada = os.environ.get("RODADA", "")
if not re.fullmatch(r"[0-9]{8}", rodada):
    sys.exit("RODADA ausente ou inválida")
VERSAO = "v8"  # v2: auditoria do M4 (A1-129); v3: auditoria do chat sobre o M4 (A1-132); v4: auditoria curta da v3 (A1-134); v5: retorno do chat e M5 (A1-138); v6: códigos para a RECORD (A1-139); v7: M5b, Figura S1 e Tabela S6 (A1-144); v8: nota da Tabela S6 sobre os denominadores da Tabela 1 (A1-147)
arq = f"logs/docx_m4_fase4_{rodada}_{VERSAO}.md"
D = "manuscrito/m4"
OUT = {k: f"{D}/{k}_{VERSAO}_{rodada}.docx" for k in ("manuscrito_cego", "pagina_titulo", "material_suplementar", "highlights")}
SUB = "Submissao"
for p in [arq, *OUT.values()]:
    if os.path.exists(p):
        sys.exit(f"já existe: {p} (nunca sobrescreve)")
# Submissao/ existe desde a Fase 0 só com o .gitkeep: aceita-se a pasta vazia (fora o .gitkeep e os ._* do disco), nunca com arquivo dentro
ja = [os.path.join(r, f) for r, _, fs in os.walk(SUB) for f in fs if f != ".gitkeep" and not f.startswith("._")] if os.path.exists(SUB) else []
if ja:
    sys.exit(f"{SUB}/ já tem arquivos (nunca sobrescreve): {ja}")
sujos = git("status", "--porcelain", "--", ESTE)
if sujos:
    sys.exit("script com mudança não commitada: não roda\n" + sujos)
head_c = git("rev-parse", "HEAD"); ramo = git("branch", "--show-current")
st = git("status", "--porcelain", "--", "artigo1_prematuridade", "codigo_compartilhado")
coord = git("log", "-1", "--format=%h %s", "--grep=^COORD:")
inicio = datetime.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
MDS = {k: f"{D}/{k}_{VERSAO}_{rodada}.md" for k in ("manuscrito", "suplemento", "pagina_titulo", "highlights", "legendas")}
TAB = {k: f"tabelas/{k}_v4_{rodada}.csv" for k in ("tab01_caracteristicas", "tab02_resultados_principais", "tab03_criterio_c1", "tabS1_sensibilidades",
                                                  "tabS2_blups", "tabS3_contrastes_concentracoes", "tabS4_nula_posthoc")}
REF = "manuscrito/referencias_m4_v4.csv"   # v3: o v2 (v1 + busca da evidência brasileira, A1-132) mais o BR-DWGD como conjunto de dados (A1-134); v4: + Szwarcwald 2019 e Bonilha 2018 (A1-144)
FIG = {f"Figure{i}.pdf": f"figuras/fig0{i}_{n}_v2_{rodada}.pdf" for i, n in ((1, "curva_combinada"), (2, "lagresposta"), (3, "curvas_concentracoes"), (4, "sensibilidades"))}
S5 = f"{D}/tabS5_decisoes_{VERSAO}_{rodada}.csv"
LOG93 = "logs/selecao_ausentes_m5b_20261020_v2.md"   # script 93 v2 (fluxograma e ausentes), rodada própria
FIGS1 = {"pdf": "figuras/figS1_fluxograma_v2_20261020.pdf", "png": "figuras/figS1_fluxograma_v2_20261020.png"}
S6 = "tabelas/tabS6_ausentes_v2_20261020.csv"
LOGS = {**{p: f"logs/manuscrito_m4_fase4_{rodada}_{VERSAO}.md" for p in [*MDS.values(), S5]}, **{p: f"logs/tabelas_m3_fase4_{rodada}_v4.md" for p in TAB.values()},
        **{p: f"logs/figuras_m3_fase4_{rodada}_v2.md" for p in FIG.values()}, **{p: LOG93 for p in [*FIGS1.values(), S6]}}
sha_ent = {}
for p, lg in LOGS.items():
    s = sha(p); sha_ent[p] = s
    if f"`{p}`: `{s}`" not in open(lg, encoding="utf-8").read():
        sys.exit(f"{p}: SHA-256 diferente do registrado em {lg}")
sha_ent[REF] = sha(REF)
refs = {r["chave"]: r for r in csv.DictReader(open(REF, encoding="utf-8"))}
ok_ref = {k for k, r in refs.items() if r.get("verificada", "").strip().lower() == "sim"}


CORRIGE = {"Fitch2025": [("Strickland M,", "Strickland MJ,")],
           "vonElm2007": [("von Elm E, Altman DG, Egger M, et al.", "von Elm E, Altman DG, Egger M, Pocock SJ, Gøtzsche PC, Vandenbroucke JP; STROBE Initiative.")],
           "Couto2026": [("Preprint posted online September 1, 2026. Research Square.", "Research Square. Preprint posted online September 1, 2026.")],
           # URL efetivamente acessado na baixada (logs/integridade.csv da base; auditoria da M4 v3, Nota 22)
           "SINASC": [("https://dadosabertos.saude.gov.br/dataset/sistema-de-informacao-sobre-nascidos-vivos-sinasc", "ftp://ftp.datasus.gov.br/dissemin/publicos/SINASC/1996_/Dados/DNRES/")],
           "SIM": [("https://dadosabertos.saude.gov.br/dataset/sim", "ftp://ftp.datasus.gov.br/dissemin/publicos/SIM/CID10/")],
           "IBGE2016": [("https://biblioteca.ibge.gov.br/visualizacao/livros/liv99700.pdf", "https://geoftp.ibge.gov.br/organizacao_do_territorio/tipologias_do_territorio/arranjos_populacionais/base_de_dados_2ed/")]}
for k, trocas in CORRIGE.items():
    for velho, novo in trocas:
        if refs[k]["citacao_ama"].count(velho) != 1:
            sys.exit(f"correção de citação não casou: {k}")
        refs[k]["citacao_ama"] = refs[k]["citacao_ama"].replace(velho, novo)


ACESSO = {"SINASC": "October 2, 2026", "SIM": "October 2, 2026", "BRDWGD": "October 2, 2026", "CAMS_EAC4": "October 5, 2026", "IBGE2016": "October 6, 2026", "AtlasDH": "October 19, 2026",
          "RCore": "October 8, 2026", "gnm": "October 8, 2026"}   # base: logs/integridade.csv (última baixada); Atlas: log do IDHM; R e gnm: congelamento da base 4443971b


def cit(k):
    """Citação AMA do CSV; dados e software ganham a data de acesso (AMA: antes do URL ou do DOI)."""
    t = refs[k]["citacao_ama"].strip()
    if refs[k].get("tipo", "").strip() in ("dataset", "software") and "Accessed" not in t:
        if k not in ACESSO:
            sys.exit(f"sem data de acesso: {k}")
        marca = next(m for m in (" http", " ftp", " doi:") if m in t)
        i = t.index(marca) if marca == " doi:" else t.rindex(marca)
        t = t[:i] + f" Accessed {ACESSO[k]}." + t[i:]
    return t

# ------------------------------------------------------------------ utilidades
CIT = re.compile(r"\[@([^\]]+)\]")


def numera(textos, ordem):
    """Atribui números às chaves na ordem da primeira citação; devolve as chaves sem verificação."""
    falta = []
    for t in textos:
        for m in CIT.finditer(t):
            for k in [x.strip().lstrip("@") for x in m.group(1).split(";")]:
                if k not in ok_ref:
                    falta.append(k)
                elif k not in ordem:
                    ordem[k] = len(ordem) + 1
    return sorted(set(falta))


def faixa(ns):
    ns = sorted(set(ns)); out = []; i = 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(str(ns[i]) if j == i else (f"{ns[i]},{ns[j]}" if j == i + 1 else f"{ns[i]}–{ns[j]}"))
        i = j + 1
    return ",".join(out)


def novo_doc(linhas_numeradas=True, espaco=2.0):
    d = Document()
    stn = d.styles["Normal"]; stn.font.name = "Arial"; stn.font.size = Pt(11)
    stn.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    stn.paragraph_format.line_spacing = espaco
    for nm in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        if nm in [s.name for s in d.styles]:
            d.styles[nm].font.name = "Arial"
    sec = d.sections[0]
    if linhas_numeradas:
        ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
        depois = [c for c in sec._sectPr if c.tag in (qn("w:pgNumType"), qn("w:cols"), qn("w:formProt"), qn("w:vAlign"), qn("w:noEndnote"), qn("w:titlePg"), qn("w:textDirection"), qn("w:bidi"), qn("w:rtlGutter"), qn("w:docGrid"))]
        depois[0].addprevious(ln) if depois else sec._sectPr.append(ln)
    cp = d.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.title = ""; cp.subject = ""; cp.keywords = ""
    ss = d.settings.element; ah = OxmlElement("w:autoHyphenation"); ah.set(qn("w:val"), "false")
    for velho in ss.findall(qn("w:autoHyphenation")):
        ss.remove(velho)
    dts = ss.find(qn("w:defaultTabStop"))
    dts.addnext(ah) if dts is not None else ss.insert(0, ah)
    return d


def par(d, texto, ordem, negrito=False, estilo=None):
    p = d.add_paragraph(style=estilo) if estilo else d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pos = 0
    for m in CIT.finditer(texto):
        antes = texto[pos:m.start()]; pontua = ""
        if antes and antes[-1] in ":;":
            antes, pontua = antes[:-1], antes[-1]
        if antes:
            r = p.add_run(antes); r.bold = negrito
        ks = [x.strip().lstrip("@") for x in m.group(1).split(";")]
        if all(k in ordem for k in ks):
            r = p.add_run(faixa([ordem[k] for k in ks])); r.font.superscript = True
        else:
            r = p.add_run("[ref]")
        if pontua:
            r = p.add_run(pontua); r.bold = negrito
        pos = m.end()
    if pos < len(texto):
        r = p.add_run(texto[pos:]); r.bold = negrito
    return p


def cab(d, texto, nivel):
    h = d.add_heading(texto, level=nivel); h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in h.runs:
        r.font.name = "Arial"


def md_para_doc(d, caminho, ordem):
    n = 0
    for linha in open(caminho, encoding="utf-8").read().split("\n"):
        if linha.startswith("### "):
            cab(d, linha[4:], 3)
        elif linha.startswith("## "):
            cab(d, linha[3:], 2)
        elif linha.startswith("# "):
            cab(d, linha[2:], 1)
        elif linha.startswith("- "):
            par(d, linha[2:], ordem, estilo="List Bullet"); n += 1
        elif linha.strip():
            par(d, linha, ordem); n += 1
    return n


def bordas_tres_linhas(t):
    tbl = t._tbl; tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), val)
        if val == "single":
            e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    for velho in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(velho)
    seguinte = [c for c in tblPr if c.tag in (qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook"), qn("w:tblCaption"), qn("w:tblDescription"))]
    seguinte[0].addprevious(b) if seguinte else tblPr.append(b)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); tb.append(e)
        tcPr.append(tb)


def tabela(d, titulo, p_csv, nota, ordem):
    par(d, titulo, ordem, negrito=True)
    rows = list(csv.reader(open(p_csv, encoding="utf-8")))
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for k, v in enumerate(r):
            c = t.cell(i, k); c.text = v
            for pp in c.paragraphs:
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT; pp.paragraph_format.line_spacing = 1.0
                for run in pp.runs:
                    run.font.size = Pt(8); run.font.bold = (i == 0)
    bordas_tres_linhas(t)
    if nota:
        par(d, nota, ordem)


def confere_tabelas(caminho):
    """Nenhuma borda vertical (insideV/left/right) ativa e nenhum sombreamento."""
    xml = Document(caminho).element.xml
    ruins = re.findall(r'<w:(insideV|left|right|start|end) w:val="(?!nil|none)[a-zA-Z]+"', xml)
    sombra = re.findall(r"<w:shd ", xml)
    return len(ruins), len(sombra)


def justificado(caminho):
    d = Document(caminho)
    al = [str(p.alignment) for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                al += [str(p.alignment) for p in c.paragraphs if p.text.strip()]
    return any("JUSTIFY" in a for a in al)


NOTAS = {
    "tabS5_decisoes": ("Supplementary Table S1. Decisions taken after the analysis plan was version-controlled",
        "Decisions are numbered as in the project's decision log; D9 and D14 concern another analysis of the project and do not apply to this article. Post hoc: chosen after the plan was version-controlled; clarification of the plan: interpretation of the prespecified text. The full log is available with the analysis code. Used results of this study: whether any estimate of this study entered the decision. MRT, minimum-risk temperature; CI, confidence interval; RR, relative risk; DLNM, distributed lag nonlinear model; IBGE, Brazilian Institute of Geography and Statistics."),
    "tab01_caracteristicas": ("Table 1. Characteristics of the 26 urban concentrations, Brazil, 2012–2024",
        "Preterm births are counts in the main analysis window (April 14, 2012, to August 6, 2024). Cesarean delivery and ≥7 prenatal visits are proportions among singleton live births at 22–42 weeks, 2012–2024, with unknown values excluded from the denominator. Temperatures are daily means of the core city, 2012–2024. MRT, minimum-risk temperature; BLUP, best linear unbiased prediction; two-letter codes are Brazilian state abbreviations. † MRT at the 99th percentile (grid boundary)."),
    "tab02_resultados_principais": ("Table 2. Pooled results of the primary analysis",
        "RR, cumulative relative risk over lags 0–7 against the pooled MRT (minimum-risk temperature), predicted at the mean of the meta-predictors; CI, confidence interval; df, degrees of freedom; BLUP, best linear unbiased prediction. Attributable fractions and numbers are reported only as the burden that would be attributable if the associations were causal; eCI, empirical 95% confidence interval (1,000 Monte Carlo samples). Forward perspective with BLUP-reduced coefficients (primary); backward perspective with first-stage coefficients (sensitivity)."),
    "tab03_criterio_c1": ("Table 3. Key sensitivity analyses and the prespecified robustness criterion",
        "Each analysis re-estimated its own pooled MRT (minimum-risk temperature). RR, cumulative relative risk over lags 0–7; CI, confidence interval; PM2.5, fine particulate matter (diameter ≤2.5 µm); γ, coefficient of the logarithm of the delivery-risk weight; (i), spontaneous preterm birth with births of unknown onset excluded from the numerator; ns, natural spline; df, degrees of freedom. A result was considered robust if the estimate had the same direction as the primary estimate and fell within its 95% CI; the 99th and 1st percentile contrasts were assessed separately."),
    "tabS1_sensibilidades": ("Supplementary Table S2. Closed list of sensitivity analyses",
        "The table lists the primary analysis and 43 other rows: 41 analyses run, the negative binomial model (not triggered), and the 2,000-event threshold (no concentration excluded); the 6 key sensitivity analyses are in Table 3. Boundary: MRT at the grid boundary on that side (contrast fixed at 1, without CI); for these analyses, RRs against the median are shown. Exposure-metric and warm-season analyses use their own temperature axes. Exploratory meta-regressions add one characteristic at a time to the mean and range of temperature. ID, analysis identifier; MRT, minimum-risk temperature; RR, relative risk; CI, confidence interval; GA, gestational age; LMP, last menstrual period; SINASC, Brazilian Live Birth Information System; PM2.5, fine particulate matter; ns, natural spline; DLNM, distributed lag nonlinear model; MHDI, municipal human development index; (i) and (ii), versions of the outcome definition (Supplementary Methods)."),
    "tabS2_blups": ("Supplementary Table S4. Concentration-specific estimates (best linear unbiased predictions)",
        "† MRT at the 99th percentile (grid boundary): the 99th percentile contrast equals 1 by construction and is shown without a CI; for these concentrations, the minimum searched between the 10th and 90th percentiles (with its bootstrap 95% CI) and RRs against the median are shown; ‡ the searched minimum fell at an end of the search range. MRT, minimum-risk temperature; RR, relative risk; CI, confidence interval; two-letter codes are Brazilian state abbreviations."),
    "tabS3_contrastes_concentracoes": ("Supplementary Table S3. Concentration-specific secondary contrasts (best linear unbiased predictions), in percentiles and °C",
        "RR, cumulative relative risk over lags 0–7 against the concentration-specific MRT (minimum-risk temperature); CI, confidence interval; two-letter codes are Brazilian state abbreviations."),
    "tabS6_ausentes": ("Supplementary Table S6. Missing or unknown values in the analytic population",
        "Totals for the 26 urban concentrations; the analytic population comprises 13,509,914 singleton live births at 22–42 weeks in the fixed cohort, April 14, 2012, to August 6, 2024. Field names are those of the SINASC files (Supplementary Methods). Missing or unknown: PARTO other than vaginal or cesarean; CONSULTAS other than codes 1–4; STTRABPART and STCESPARTO other than yes or no (STCESPARTO among cesarean deliveries only); PESO missing or outside 200–7,000 g; DTULTMENST missing, not a valid date, after the birth, or more than 350 days before it. Gestational age is complete by definition of the population; temperature is counted in concentration-days of the analysis window. Proportions in Table 1 use singleton live births at 22–42 weeks, 2012–2024, and its temperatures cover 2012–2024. Counts from 1 to 9 would be shown as <10. SINASC, Brazilian Live Birth Information System."),
    "tabS4_nula_posthoc": ("Supplementary Table S5. Post hoc analysis: distribution under no association in any urban concentration (1,000 replicates)",
        "Replicates drew first-stage reduced coefficients from multivariate normal distributions with mean zero and the estimated first-stage covariance matrices, and repeated the second-stage meta-analysis, the MRT search, and the forward attributable-fraction computation; no replicate failed. MRT, minimum-risk temperature; RR, relative risk."),
}

FIGS1_LEGENDA = ("Exclusions are shown in the order in which they were applied to live births to residents of the 277 municipalities of the 26 urban concentrations in the main analysis window. "
                 "Gestational age is the SINASC field in completed weeks; values outside 22–44 weeks were treated as missing, and births at 43–44 weeks were outside the study population. "
                 "The estimated pregnancy start date was the date of birth minus (7 × gestational age + 3) days; births were restricted to a fixed cohort of pregnancies to avoid fixed-cohort bias. "
                 "Counts are totals for the 26 urban concentrations. SINASC, Brazilian Live Birth Information System.")

# ------------------------------------------------------------------ manuscrito cego
ordem = {}
txt_m = open(MDS["manuscrito"], encoding="utf-8").read()
txt_l = open(MDS["legendas"], encoding="utf-8").read()
notas_txt = " ".join(n for _, n in (NOTAS[k] for k in ("tab01_caracteristicas", "tab02_resultados_principais", "tab03_criterio_c1")))
falta_m = numera([txt_m, notas_txt, txt_l], ordem)
dm = novo_doc()
n_par = md_para_doc(dm, MDS["manuscrito"], ordem)
dm.add_page_break(); cab(dm, "References", 1)
for k, n in sorted(ordem.items(), key=lambda kv: kv[1]):
    par(dm, f"{n}. {cit(k)}", {})
dm.add_page_break(); cab(dm, "Tables", 1)
for k in ("tab01_caracteristicas", "tab02_resultados_principais", "tab03_criterio_c1"):
    tit, nota = NOTAS[k]; tabela(dm, tit, TAB[k], nota, ordem); dm.add_paragraph()
dm.add_page_break(); md_para_doc(dm, MDS["legendas"], ordem)
dm.save(OUT["manuscrito_cego"])

# ------------------------------------------------------------------ material suplementar (referências próprias)
ordem_s = {}
txt_s = open(MDS["suplemento"], encoding="utf-8").read()
falta_s = numera([txt_s], ordem_s)
ds = novo_doc(linhas_numeradas=False, espaco=1.15)
md_para_doc(ds, MDS["suplemento"], ordem_s)
ds.add_page_break(); par(ds, "Supplementary Figure S1. Selection of the analytic population", ordem_s, negrito=True)
ds.add_picture(FIGS1["png"], width=Inches(6.0))
par(ds, FIGS1_LEGENDA, ordem_s)
ds.add_page_break(); tit, nota = NOTAS["tabS5_decisoes"]; tabela(ds, tit, S5, nota, ordem_s)
for k in ("tabS1_sensibilidades", "tabS3_contrastes_concentracoes", "tabS2_blups", "tabS4_nula_posthoc"):   # S2, S3, S4, S5
    ds.add_page_break(); tit, nota = NOTAS[k]; tabela(ds, tit, TAB[k], nota, ordem_s)
ds.add_page_break(); tit, nota = NOTAS["tabS6_ausentes"]; tabela(ds, tit, S6, nota, ordem_s)
if ordem_s:
    ds.add_page_break(); cab(ds, "Supplementary references", 1)
    for k, n in sorted(ordem_s.items(), key=lambda kv: kv[1]):
        par(ds, f"{n}. {cit(k)}", {})
ds.save(OUT["material_suplementar"])

# ------------------------------------------------------------------ página de título e highlights
dt = novo_doc(linhas_numeradas=False, espaco=1.15); md_para_doc(dt, MDS["pagina_titulo"], {}); dt.save(OUT["pagina_titulo"])
dh = novo_doc(linhas_numeradas=False, espaco=1.15); md_para_doc(dh, MDS["highlights"], {}); dh.save(OUT["highlights"])

# ------------------------------------------------------------------ conferências
conf = {}
for k, p in OUT.items():
    nv, ns = confere_tabelas(p); conf[k] = (nv, ns, justificado(p))
xml_m = Document(OUT["manuscrito_cego"]).element.xml
linhas_ok = "w:lnNumType" in Document(OUT["manuscrito_cego"]).sections[0]._sectPr.xml
CEGO = ["Guilherme", "Coelho", "UNICAMP", "Unicamp", "CliMaterna", "GitHub", "guilcoel", "/Volumes", "artigos_perinatais"]
texto_m = "\n".join(p.text for p in Document(OUT["manuscrito_cego"]).paragraphs)
vaz = [w for w in CEGO if w in texto_m]
falhas = [k for k, (nv, ns, ju) in conf.items() if nv or ns or ju] + (["numeração de linhas"] if not linhas_ok else []) + (["duplo-cego: " + ", ".join(vaz)] if vaz else [])
if not falhas:
    os.makedirs(f"{SUB}/figuras", exist_ok=True)
    for nome, p in (("manuscrito_cego.docx", OUT["manuscrito_cego"]), ("pagina_de_titulo.docx", OUT["pagina_titulo"]), ("material_suplementar.docx", OUT["material_suplementar"]), ("highlights.docx", OUT["highlights"])):
        shutil.copyfile(p, f"{SUB}/{nome}")
    for nome, p in FIG.items():
        shutil.copyfile(p, f"{SUB}/figuras/{nome}")
    shutil.copyfile(FIGS1["pdf"], f"{SUB}/figuras/FigureS1.pdf")
sub_lista = sorted(os.path.relpath(os.path.join(r, f), SUB) for r, _, fs in os.walk(SUB) for f in fs if f != ".gitkeep" and not f.startswith("._")) if os.path.exists(SUB) else []
fim = datetime.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
md = ["# Fase 4, marco M4: .docx da submissão", "",
      f"- **Gerado por:** `00_scripts/91_docx_m4_fase4.py` (SHA-256 `{sha(os.path.join(REPO, ESTE))}`), início em {inicio}, fim em {fim}. Rodada `{rodada}`.",
      f"- **Cabeçalho (protocolo v1.1, 3.8):** `HEAD` = `{head_c}` (ramo `{ramo}`); último commit da `COORD` no `HEAD`: `{coord}`; `git status --porcelain -- artigo1_prematuridade codigo_compartilhado`: {('`' + '`; `'.join(st.splitlines()) + '`') if st else 'vazio'}.",
      "- **Entradas (SHA-256 = registrados; o CSV de referências é do subagente):**"] + [f"  - `{p}`: `{s}`;" for p, s in sha_ent.items()] + [
      f"- **Referências no manuscrito:** {len(ordem)} numeradas na ordem da primeira citação; chaves sem referência verificada (viram [ref]): {', '.join(falta_m) if falta_m else 'nenhuma'}.",
      f"- **Referências do suplemento:** {len(ordem_s)}; sem verificação: {', '.join(falta_s) if falta_s else 'nenhuma'}.",
      "- **Conferências** (bordas verticais ativas; sombreamento; texto justificado): " + "; ".join(f"{k}: {nv}, {ns}, {'sim' if ju else 'não'}" for k, (nv, ns, ju) in conf.items()) + ".",
      f"- **Numeração de linhas contínua no manuscrito:** {'sim' if linhas_ok else 'não'}. **Hifenização automática:** desligada. **Duplo-cego** (nomes, instituição, caminhos): {'limpo' if not vaz else 'FALHA: ' + ', '.join(vaz)}.",
      f"- **Pasta `{SUB}/`:** {', '.join(sub_lista) if sub_lista else 'não criada (falha)'}.", "",
      "## Saídas (fora do git)", ""] + [f"- `{p}`: `{sha(p)}`;" for p in OUT.values()] + ["", f"- **Resultado:** {'FALHA: ' + '; '.join(falhas) if falhas else 'ok'}.", ""]
open(arq, "w", encoding="utf-8").write("\n".join(md))
print("gravado:", arq, "| falhas:", falhas, "| refs:", len(ordem), "| sem fonte:", falta_m)
if falhas:
    sys.exit(1)
