# 95_suplemento_entrega.py — material suplementar do A1 no padrão da entrega (decisões de Guilherme de 21/10; A1-153): um .docx com folha de rosto
# e sumário, Supplementary Methods, Supplementary Figure S1, Supplementary Tables S1–S6 (na ordem da primeira citação no texto principal) e
# Supplementary References. Só .docx (decisão de Guilherme). Fora do git: Entrega/suplemento/.
# Entradas (SHA-256 conferido contra os logs de origem): suplemento_v8_20261019.md e tabS5_decisoes_v8_20261019.csv (log do 90 v8); tabelas
#   tabS1–tabS4 v4 (log do 87 v4); tabS6_ausentes_v2 e figS1_fluxograma_v2 (log do 93 v2); manuscrito/referencias_m4_v4.csv.
# Renumeração (A1-153), por mapa único: decisões S1; lista fechada S2; dados ausentes S3 (era S6); contrastes por concentração S4 (era S3);
#   BLUPs S5 (era S4); distribuição post hoc S6 (era S5).
# Padrão da entrega (A1-151): Arial 11, espaço 1,15, número de página, títulos pretos, P em itálico, sem justificar nem hifenizar; tabelas de três
#   linhas com título acima e nota abaixo; tabelas largas em páginas deitadas; referências com [dataset] antes e DOI por último; frases prescritas intactas.
# v2 (A1-154): tabelas com cabeçalho repetido, linhas sem quebra, espaço zero e larguras proporcionais, em 8 pt; notas da S1, S2 e S3 completas; df na
#   lista de abreviações; sumário com os nomes das subseções; metadados com a data da rodada; saída da v1 na lixeira.
# Funções de montagem copiadas do 94 v4. Log logs/suplemento_entrega_<RODADA>_<VERSAO>.md.
# Uso (de artigo1_prematuridade/):  RODADA=AAAAMMDD python3 -I 00_scripts/95_suplemento_entrega.py
import csv, datetime, hashlib, os, re, subprocess, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

REPO = "/Volumes/Forja/Doc/artigos_perinatais_clima"
ESTE = "artigo1_prematuridade/00_scripts/95_suplemento_entrega.py"
VERSAO = "v2"
TITULO = "Short-term ambient temperature exposure and preterm birth in 26 Brazilian urban concentrations"


def git(*a):
    return subprocess.run(["git", "-C", REPO, *a], capture_output=True, text=True, check=True).stdout.strip()


def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


rodada = os.environ.get("RODADA", "")
if not re.fullmatch(r"[0-9]{8}", rodada):
    sys.exit("RODADA ausente ou inválida")
ENT = "20261019"
arq = f"logs/suplemento_entrega_{rodada}_{VERSAO}.md"
D = "Entrega/suplemento"
OUT = f"{D}/A1_material_suplementar_{VERSAO}_{rodada}.docx"
for p in (arq, OUT):
    if os.path.exists(p):
        sys.exit(f"já existe: {p} (nunca sobrescreve)")
ANTERIOR = {"Entrega/suplemento/A1_material_suplementar_v1_20261021.docx": "f442196ab05e5d6a250dfde1056bcb4fad9ef3a38fd7ed232e672023f620afa6"}
for p in ANTERIOR:
    if os.path.exists(p):
        sys.exit(f"saída da v1 ainda no lugar (mova para a lixeira antes): {p}")
sujos = git("status", "--porcelain", "--", ESTE)
if sujos:
    sys.exit("script com mudança não commitada: não roda\n" + sujos)
head_c = git("rev-parse", "HEAD"); ramo = git("branch", "--show-current")
st = git("status", "--porcelain", "--", "artigo1_prematuridade", "codigo_compartilhado")
coord = git("log", "-1", "--format=%h %s", "--grep=^COORD:")
inicio = datetime.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")

# ------------------------------------------------------------------ entradas (SHA-256 = registrados)
SUP = f"manuscrito/m4/suplemento_v8_{ENT}.md"; PT = f"manuscrito/m4/pagina_titulo_v8_{ENT}.md"
TAB = {   # número novo → (CSV, log de origem)
    "S1": (f"manuscrito/m4/tabS5_decisoes_v8_{ENT}.csv", f"logs/manuscrito_m4_fase4_{ENT}_v8.md"),
    "S2": (f"tabelas/tabS1_sensibilidades_v4_{ENT}.csv", f"logs/tabelas_m3_fase4_{ENT}_v4.md"),
    "S3": ("tabelas/tabS6_ausentes_v2_20261020.csv", "logs/selecao_ausentes_m5b_20261020_v2.md"),
    "S4": (f"tabelas/tabS3_contrastes_concentracoes_v4_{ENT}.csv", f"logs/tabelas_m3_fase4_{ENT}_v4.md"),
    "S5": (f"tabelas/tabS2_blups_v4_{ENT}.csv", f"logs/tabelas_m3_fase4_{ENT}_v4.md"),
    "S6": (f"tabelas/tabS4_nula_posthoc_v4_{ENT}.csv", f"logs/tabelas_m3_fase4_{ENT}_v4.md"),
}
FIGS1 = ("figuras/figS1_fluxograma_v2_20261020.png", "logs/selecao_ausentes_m5b_20261020_v2.md")
REF = "manuscrito/referencias_m4_v4.csv"
LOGS = {SUP: f"logs/manuscrito_m4_fase4_{ENT}_v8.md", PT: f"logs/manuscrito_m4_fase4_{ENT}_v8.md", FIGS1[0]: FIGS1[1], **{c: lg for c, lg in TAB.values()}}
sha_ent = {}
for p, lg in LOGS.items():
    s = sha(p); sha_ent[p] = s
    if f"`{p}`: `{s}`" not in open(lg, encoding="utf-8").read():
        sys.exit(f"{p}: SHA-256 diferente do registrado em {lg}")
sha_ent[REF] = sha(REF)
if TITULO not in open(PT, encoding="utf-8").read():
    sys.exit("título diferente do da página de título v8")
refs = {r["chave"]: r for r in csv.DictReader(open(REF, encoding="utf-8"))}
ok_ref = {k for k, r in refs.items() if r.get("verificada", "").strip().lower() == "sim"}

# ------------------------------------------------------------------ referências (como no 94 v4)
CORRIGE = {"Fitch2025": [("Strickland M,", "Strickland MJ,")],
           "vonElm2007": [("von Elm E, Altman DG, Egger M, et al.", "von Elm E, Altman DG, Egger M, Pocock SJ, Gøtzsche PC, Vandenbroucke JP; STROBE Initiative.")],
           "Couto2026": [("Preprint posted online September 1, 2026. Research Square.", "Research Square. Preprint posted online September 1, 2026.")],
           "SINASC": [("https://dadosabertos.saude.gov.br/dataset/sistema-de-informacao-sobre-nascidos-vivos-sinasc", "ftp://ftp.datasus.gov.br/dissemin/publicos/SINASC/1996_/Dados/DNRES/"), ("[dataset]. DATASUS.", "[dataset]. DATASUS. Data years 2011–2024.")],
           "SIM": [("https://dadosabertos.saude.gov.br/dataset/sim", "ftp://ftp.datasus.gov.br/dissemin/publicos/SIM/CID10/"), ("[dataset]. DATASUS.", "[dataset]. DATASUS. Data years 2011–2024.")],
           "IBGE2016": [("https://biblioteca.ibge.gov.br/visualizacao/livros/liv99700.pdf", "https://geoftp.ibge.gov.br/organizacao_do_territorio/tipologias_do_territorio/arranjos_populacionais/base_de_dados_2ed/")],
           "gnm": [(". https://CRAN.R-project.org/package=gnm", "")]}
for k, trocas in CORRIGE.items():
    for velho, novo in trocas:
        if refs[k]["citacao_ama"].count(velho) != 1:
            sys.exit(f"correção de citação não casou: {k}")
        refs[k]["citacao_ama"] = refs[k]["citacao_ama"].replace(velho, novo)
ACESSO = {"SINASC": "October 2, 2026", "SIM": "October 2, 2026", "BRDWGD": "October 2, 2026", "CAMS_EAC4": "October 5, 2026", "IBGE2016": "October 6, 2026", "AtlasDH": "October 19, 2026",
          "RCore": "October 8, 2026", "gnm": "October 8, 2026"}


def cit(k):
    t = refs[k]["citacao_ama"].strip(); tipo = refs[k].get("tipo", "").strip()
    if tipo in ("dataset", "software") and "Accessed" not in t:
        if k not in ACESSO:
            sys.exit(f"sem data de acesso: {k}")
        marca = next(m for m in (" http", " ftp", " doi:") if m in t)
        i = t.index(marca) if marca == " doi:" else t.rindex(marca)
        t = t[:i] + f" Accessed {ACESSO[k]}." + t[i:]
    if tipo == "dataset":
        t = "[dataset] " + t.replace(" [dataset]", "")
    return t


CIT = re.compile(r"\[@([^\]]+)\]")


def numera(textos, ordem):
    falta = []
    for t in textos:
        for m in CIT.finditer(t):
            for k in [x.strip().lstrip("@") for x in m.group(1).split(";")]:
                if k not in ok_ref:
                    falta.append(k)
                elif k not in ordem:
                    ordem[k] = len(ordem) + 1
    return sorted(set(falta))


def faixa(ns):
    ns = sorted(set(ns)); out = []; i = 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(str(ns[i]) if j == i else (f"{ns[i]},{ns[j]}" if j == i + 1 else f"{ns[i]}–{ns[j]}"))
        i = j + 1
    return ",".join(out)


def novo_doc():
    d = Document()
    stn = d.styles["Normal"]; stn.font.name = "Arial"; stn.font.size = Pt(11)
    stn.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    stn.paragraph_format.line_spacing = 1.15
    for nm in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        if nm in [s.name for s in d.styles]:
            d.styles[nm].font.name = "Arial"; d.styles[nm].font.color.rgb = RGBColor(0, 0, 0)
    rodape(d.sections[0])
    cp = d.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.title = ""; cp.subject = ""; cp.keywords = ""
    agora_dt = datetime.datetime.now().replace(microsecond=0); cp.created = agora_dt; cp.modified = agora_dt; cp.revision = 1   # A1-154: sem as datas do modelo
    ss = d.settings.element; ah = OxmlElement("w:autoHyphenation"); ah.set(qn("w:val"), "false")
    for velho in ss.findall(qn("w:autoHyphenation")):
        ss.remove(velho)
    dts = ss.find(qn("w:defaultTabStop"))
    dts.addnext(ah) if dts is not None else ss.insert(0, ah)
    return d


def rodape(sec):
    """Número de página (campo PAGE), à direita."""
    fp = sec.footer.paragraphs[0]
    if "PAGE" in fp._p.xml:
        return
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); fld.append(rr); fp._p.append(fld)


def nova_secao(d, deitada):
    """Quebra de seção com orientação própria (tabelas largas em página deitada)."""
    sec = d.add_section()
    larg, alt = sorted((sec.page_width, sec.page_height))
    sec.orientation = WD_ORIENT.LANDSCAPE if deitada else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (alt, larg) if deitada else (larg, alt)
    return sec


P_ITAL = re.compile(r"(?<![A-Za-z0-9])P(?=\s?[=<>]|;|$)")


def texto_italico(p, texto, negrito=False, tamanho=None):
    pos = 0; runs = []
    for m in P_ITAL.finditer(texto):
        if m.start() > pos:
            runs.append((texto[pos:m.start()], False))
        runs.append(("P", True)); pos = m.end()
    if pos < len(texto):
        runs.append((texto[pos:], False))
    for t, it in runs:
        r = p.add_run(t); r.bold = negrito; r.italic = it or None
        if tamanho:
            r.font.size = Pt(tamanho)


def par(d, texto, ordem, negrito=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pos = 0
    for m in CIT.finditer(texto):
        antes = texto[pos:m.start()]; pontua = ""
        if antes and antes[-1] in ":;":
            antes, pontua = antes[:-1], antes[-1]
        if antes:
            texto_italico(p, antes, negrito)
        ks = [x.strip().lstrip("@") for x in m.group(1).split(";")]
        if all(k in ordem for k in ks):
            r = p.add_run(faixa([ordem[k] for k in ks])); r.font.superscript = True
        else:
            r = p.add_run("[ref]")
        if pontua:
            r = p.add_run(pontua); r.bold = negrito
        pos = m.end()
    if pos < len(texto):
        texto_italico(p, texto[pos:], negrito)
    return p


def cab(d, texto, nivel):
    h = d.add_heading(texto, level=nivel); h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in h.runs:
        r.font.name = "Arial"


def bordas_tres_linhas(t):
    tbl = t._tbl; tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), val)
        if val == "single":
            e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    for velho in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(velho)
    seguinte = [c for c in tblPr if c.tag in (qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook"), qn("w:tblCaption"), qn("w:tblDescription"))]
    seguinte[0].addprevious(b) if seguinte else tblPr.append(b)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); tb.append(e)
        tcPr.append(tb)


def formata_tabela(d, t, rows):
    """A1-154: cabeçalho repetido em cada página (tblHeader), linhas sem quebra (cantSplit) e larguras proporcionais ao texto na largura útil da seção."""
    sec = d.sections[-1]
    util = sec.page_width - sec.left_margin - sec.right_margin   # EMU
    pesos = []
    for j in range(len(rows[0])):
        col = [r[j] for r in rows]
        palavra = max(len(w) for c in col for w in (c.split() or [""]))
        pesos.append(max(min(max(len(c) for c in col), 45), palavra, 4))
    larg = [int(util * w / sum(pesos)) for w in pesos]
    t.autofit = False
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(larg[j] // 635))   # EMU → twips
    for i, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            c.width = Emu(larg[j])
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trPr.append(OxmlElement("w:tblHeader"))


def tabela(d, titulo, p_csv, nota, ordem, tamanho=8):
    par(d, titulo, ordem, negrito=True)
    rows = list(csv.reader(open(p_csv, encoding="utf-8")))
    rows = [[v.replace("% [eCI]", "% (eCI)") for v in r] for r in rows]
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for k, v in enumerate(r):
            c = t.cell(i, k); pp = c.paragraphs[0]
            texto_italico(pp, v, negrito=(i == 0), tamanho=tamanho)
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT; pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_before = Pt(0); pp.paragraph_format.space_after = Pt(0)   # A1-154
    bordas_tres_linhas(t); formata_tabela(d, t, rows)
    if nota:
        par(d, nota, ordem)
    return len(rows) - 1, len(rows[0])


# ------------------------------------------------------------------ texto do suplemento (90 v8) com o padrão da entrega e a renumeração
MAPA = {"S3": "S4", "S4": "S5", "S5": "S6", "S6": "S3"}   # A1-153: número antigo → novo (S1 e S2 não mudam)


def renumera(t):
    """Troca de remissões por mapa único (sem encadear trocas)."""
    return re.sub(r"Supplementary Table (S[3-6])\b", lambda m: "Supplementary Table " + MAPA[m.group(1)], t)


L = open(SUP, encoding="utf-8").read().split("\n")
cit_antes = [m.group(0) for l in L for m in re.finditer(r"Supplementary (?:Table|Figure) S\d", l)]
L = [renumera(l) for l in L]
cit_depois = [m.group(0) for l in L for m in re.finditer(r"Supplementary (?:Table|Figure) S\d", l)]
TROCAS = [
    ("Abbreviations: BLUP, best linear unbiased prediction; CI, confidence interval; MCC,",
     "Abbreviations: BLUP, best linear unbiased prediction; CI, confidence interval; DATASUS, Department of Informatics of the Brazilian Unified Health System; df, degrees of freedom; MCC,"),
    ("Preterm subtypes used an offset restricted to the gestational weeks of the subtype, with the delivery-risk weights of the primary analysis;",
     "Preterm subtypes used an offset restricted to the gestational weeks of the subtype, with the delivery-risk weights of the primary analysis, and the coefficient of the logarithm of the delivery-risk weight was fixed at 1 for extremely and very preterm birth by the prespecified rule (approximate standard error of 1 or greater);"),
]
for velho, novo in TROCAS:
    if sum(l.count(velho) for l in L) != 1:
        sys.exit(f"troca não casou uma vez: {velho[:50]}")
    L = [l.replace(velho, novo) for l in L]
# frases prescritas do suplemento (lit do 90: SEEDS, POSTHOC, D8E, SOFT) intactas
txt_v8 = open(SUP, encoding="utf-8").read(); txt_novo = "\n".join(L)
PRESCRITAS = ["Monte Carlo samples for the empirical", "Post hoc analyses (not prespecified)", "When a concentration-specific MRT fell at the 99th percentile", "Analyses were performed in R version 4.5.1 using the packages"]
for frase in PRESCRITAS:
    ini = txt_v8.find(frase)
    if ini < 0:
        sys.exit(f"frase prescrita ausente no suplemento v8: {frase}")
    fim_f = txt_v8.find(".", ini + len(frase))
    if txt_v8[ini:fim_f] not in txt_novo:
        sys.exit(f"frase prescrita alterada: {frase}")

NOTAS = {
    "S1": ("Supplementary Table S1. Decisions taken after the analysis plan was version-controlled",
           "Decisions are numbered as in the project's decision log; D9 and D14 concern another analysis of the project and do not apply to this article. Post hoc: chosen after the plan was version-controlled; clarification of the plan: interpretation of the prespecified text. The row \"Complement\" covers the operating rules labeled D4. The full log is kept with the analysis code. Used results of this study: whether any estimate of this study entered the decision. BLUP, best linear unbiased prediction; MRT, minimum-risk temperature; CI, confidence interval; RR, relative risk; DLNM, distributed lag nonlinear model; IBGE, Brazilian Institute of Geography and Statistics.", False),
    "S2": ("Supplementary Table S2. Closed list of sensitivity analyses",
           "The table lists the primary analysis and 43 other rows: 41 analyses run, the negative binomial model (not triggered), and the 2,000-event threshold (no concentration excluded); the 6 key sensitivity analyses are in Table 3. Boundary: MRT at the grid boundary on that side (contrast fixed at 1, without CI); for these analyses, RRs against the median are shown. Exposure-metric and warm-season analyses use their own temperature axes. Exploratory meta-regressions add one characteristic at a time to the mean and range of temperature. In the 2011 cohort analyses, preterm birth was defined by the categorical gestational-age field, without offset or fixed cohort, in 2011–2024 (arm A) and 2012–2024 (arm B); the zero-missing series treats exact zeros of PM2.5 as missing; t+1 to t+7, temperature on the 7 days after the day of the count (negative control exposure); possible-event window, January 4, 2012, to November 15, 2024; early-term analyses use their own window, and the intersection window is its overlap with the main window (April 28, 2012, to August 6, 2024). ID, analysis identifier; MRT, minimum-risk temperature; RR, relative risk; CI, confidence interval; GA, gestational age; LMP, last menstrual period; SINASC, Brazilian Live Birth Information System; PM2.5, fine particulate matter; ns, natural spline; df, degrees of freedom; p10, p50, p75, and p90, 10th, 50th, 75th, and 90th percentiles; wk, weeks; DLNM, distributed lag nonlinear model; MHDI, municipal human development index; (i) and (ii), versions of the outcome definition (Supplementary Methods).", True),
    "S3": ("Supplementary Table S3. Missing or unknown values in the analytic population",
           "Totals for the 26 urban concentrations; the analytic population comprises 13,509,914 singleton live births at 22–42 weeks in the fixed cohort, April 14, 2012, to August 6, 2024. Field names are those of the SINASC files (Supplementary Methods). Missing or unknown: PARTO other than vaginal or cesarean; CONSULTAS (prenatal visits: 1, none; 2, 1–3; 3, 4–6; 4, 7 or more; 9, unknown) other than codes 1–4; STTRABPART and STCESPARTO other than yes or no (STCESPARTO among cesarean deliveries only); PESO missing or outside 200–7,000 g; DTULTMENST missing, not a valid date, after the birth, or more than 350 days before it. Gestational age is complete by definition of the population; temperature is counted in concentration-days of the analysis window. Proportions in Table 1 use singleton live births at 22–42 weeks, 2012–2024, and its temperatures cover 2012–2024. Counts from 1 to 9 would be shown as <10. SINASC, Brazilian Live Birth Information System.", False),
    "S4": ("Supplementary Table S4. Concentration-specific secondary contrasts (best linear unbiased predictions), in percentiles and °C",
           "RR, cumulative relative risk over lags 0–7 against the concentration-specific MRT (minimum-risk temperature); CI, confidence interval; two-letter codes are Brazilian state abbreviations.", True),
    "S5": ("Supplementary Table S5. Concentration-specific estimates (best linear unbiased predictions)",
           "† MRT at the 99th percentile (grid boundary): the 99th percentile contrast equals 1 by construction and is shown without a CI; for these concentrations, the minimum searched between the 10th and 90th percentiles (with its bootstrap 95% CI) and RRs against the median are shown; ‡ the searched minimum fell at an end of the search range. MRT, minimum-risk temperature; RR, relative risk; CI, confidence interval; two-letter codes are Brazilian state abbreviations.", True),
    "S6": ("Supplementary Table S6. Post hoc analysis: distribution under no association in any urban concentration (1,000 replicates)",
           "Replicates drew first-stage reduced coefficients from multivariate normal distributions with mean zero and the estimated first-stage covariance matrices, and repeated the second-stage meta-analysis, the MRT search, and the forward attributable-fraction computation; no replicate failed. MRT, minimum-risk temperature; RR, relative risk.", False),
}
FIG_TIT = "Supplementary Figure S1. Selection of the analytic population"
FIG_LEG = ("Exclusions are shown in the order in which they were applied to live births to residents of the 277 municipalities of the 26 urban concentrations in the main analysis window. "
           "Gestational age is the SINASC field in completed weeks; values outside 22–44 weeks were treated as missing, and births at 43–44 weeks were outside the study population. "
           "The estimated pregnancy start date was the date of birth minus (7 × gestational age + 3) days; births were restricted to a fixed cohort of pregnancies to avoid fixed-cohort bias. "
           "Counts are totals for the 26 urban concentrations. SINASC, Brazilian Live Birth Information System.")

# ------------------------------------------------------------------ documento
corpo = [l for l in L if not l.startswith("# ")]   # o título de nível 1 vira a folha de rosto
ordem = {}
falta = numera(["\n".join(corpo), FIG_LEG, " ".join(n for _, n, _ in NOTAS.values())], ordem)
if falta:
    sys.exit(f"chaves sem referência verificada: {falta}")
d = novo_doc()
# folha de rosto e sumário
par(d, TITULO, {}, negrito=True)
par(d, "Supplementary Material", {}, negrito=True)
par(d, "Contents", {}, negrito=True)
SUMARIO = [("Supplementary Methods", "; ".join(l[4:] for l in L if l.startswith("### ")))]   # v2: os nomes das subseções
SUMARIO += [(FIG_TIT.split(". ", 1)[0], FIG_TIT.split(". ", 1)[1])] + [(NOTAS[k][0].split(". ", 1)[0], NOTAS[k][0].split(". ", 1)[1]) for k in ("S1", "S2", "S3", "S4", "S5", "S6")]
SUMARIO += [("Supplementary References", "References cited in the Supplementary Material")]
for item, leg in SUMARIO:
    par(d, f"{item}: {leg}", {})
d.add_page_break()
n_par = 0
for l in corpo:
    if l.startswith("## "):
        cab(d, l[3:], 1)
    elif l.startswith("### "):
        cab(d, l[4:], 2)
    elif l.strip():
        par(d, l, ordem); n_par += 1
# Figura S1
nova_secao(d, False)
par(d, FIG_TIT, ordem, negrito=True)
d.add_picture(FIGS1[0], width=Inches(6.0))
par(d, FIG_LEG, ordem)
# Tabelas S1–S6
dims = {}
for k in ("S1", "S2", "S3", "S4", "S5", "S6"):
    tit, nota, deitada = NOTAS[k]
    nova_secao(d, deitada)
    dims[k] = tabela(d, tit, TAB[k][0], nota, ordem, tamanho=8)   # v2: 8 pt em todas
# referências
nova_secao(d, False)
cab(d, "Supplementary References", 1)
for k, n in sorted(ordem.items(), key=lambda kv: kv[1]):
    par(d, f"{n}. {cit(k)}", {})
for sec in d.sections:
    rodape(sec)
os.makedirs(D, exist_ok=True)
d.save(OUT)

# ------------------------------------------------------------------ conferências
doc = Document(OUT)
xml = doc.element.xml
ruins = re.findall(r'<w:(insideV|left|right|start|end) w:val="(?!nil|none)[a-zA-Z]+"', xml); sombra = re.findall(r"<w:shd ", xml)
al = [str(p.alignment) for p in doc.paragraphs if p.text.strip()] + [str(p.alignment) for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs if p.text.strip()]
justif = any("JUSTIFY" in a for a in al)
texto_doc = "\n".join(p.text for p in doc.paragraphs)
falhas = []
if ruins or sombra: falhas.append("bordas verticais ou sombreamento")
if justif: falhas.append("texto justificado")
if "[ref]" in texto_doc: falhas.append("[ref] no texto")
if len(doc.tables) != 6 or len(doc.inline_shapes) != 1: falhas.append("tabelas ou figura")
tit_tab = [p.text.split(".")[0] for p in doc.paragraphs if re.match(r"^Supplementary Table S\d\. ", p.text)]
if tit_tab != [f"Supplementary Table S{i}" for i in range(1, 7)]: falhas.append(f"ordem dos títulos das tabelas: {tit_tab}")
# remissões: só S1–S6 e a Figura S1; a S3 nova (ausentes) citada no suplemento; nenhuma remissão antiga sobrou
citadas = sorted(set(re.findall(r"Supplementary (?:Table S\d|Figure S1)", texto_doc)))
if any(int(x[-1]) > 6 for x in citadas if "Table" in x): falhas.append("remissão a tabela inexistente")
if "Supplementary Table S3" not in "\n".join(L): falhas.append("dados ausentes (S3) sem remissão no texto do suplemento")
ponto = [k for k in ordem if re.search(r"(doi:\S+|https?://\S+|ftp://\S+)\.$", cit(k))]
if ponto: falhas.append("referência termina em ponto depois de DOI ou URL: " + ", ".join(ponto))
if not all(cit(k).startswith("[dataset] ") for k in ordem if refs[k].get("tipo", "").strip() == "dataset"): falhas.append("[dataset]")
vaz = [w for w in ("[blinded for review]", "[anonymous review link]", "/Volumes", "/Users", "Unicamp", "UNICAMP", "REDU", "GitHub", "Guilherme") if w in texto_doc]
if vaz: falhas.append("marcação, caminho ou identificação no texto: " + ", ".join(vaz))
orient = [("deitada" if s.orientation == WD_ORIENT.LANDSCAPE else "em pé") for s in doc.sections]
rodape_ok = all("PAGE" in s.footer._element.xml for s in doc.sections)
if not rodape_ok: falhas.append("número de página")
n_pal = len(re.sub(r"\[@[^\]]*\]", "", "\n".join(l for l in corpo if l.strip() and not l.startswith("#"))).split())

# ------------------------------------------------------------------ log
fim = datetime.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
md = ["# Material suplementar do A1 no padrão da entrega", "",
      f"- **Gerado por:** `00_scripts/95_suplemento_entrega.py` (SHA-256 `{sha(os.path.join(REPO, ESTE))}`), início em {inicio}, fim em {fim}. Rodada `{rodada}`. Decisões: A1-153.",
      f"- **Cabeçalho (protocolo v1.1, 3.8):** `HEAD` = `{head_c}` (ramo `{ramo}`); último commit da `COORD` no `HEAD`: `{coord}`; `git status --porcelain -- artigo1_prematuridade codigo_compartilhado`: {('`' + '`; `'.join(st.splitlines()) + '`') if st else 'vazio'}.",
      "- **Entradas (SHA-256 = registrados nos logs de origem):**"] + [f"  - `{p}`: `{s}`;" for p, s in sha_ent.items()] + [
      f"- **Conteúdo:** folha de rosto com sumário ({len(SUMARIO)} itens, cada um com legenda curta); Supplementary Methods com {n_par} parágrafos ({n_pal} palavras); Supplementary Figure S1; Supplementary Tables S1–S6 ({'; '.join(f'{k}: {r} × {c}' for k, (r, c) in dims.items())}); {len(ordem)} referências próprias.",
      f"- **Renumeração (A1-153):** mapa {MAPA}; remissões no texto do suplemento antes: {cit_antes}; depois: {cit_depois}.",
      "- **Trocas no texto (padrão da entrega):**"] + [f"  - \"{a[:100]}…\" → \"{b[:160]}…\";" for a, b in TROCAS] + [
      "- **Notas:** S1 com \"kept with the analysis code\", BLUP e a linha \"Complement\" (regras D4); S2 com df, p10–p90, wk, braços A e B, série com zeros, t+1 a t+7 e as janelas de evento possível e de interseção; S3 com os códigos de CONSULTAS; títulos com os números novos.",
      "- **Trocas previstas na A1-153 que não ocorrem no suplemento v8:** \"pooled temperature axis\" e \"coprimary\" (o texto do suplemento não tinha \"common percentile scale\" nem \"co-primary\").",
      f"- **Cabeçalho repetido, linhas sem quebra e larguras proporcionais (A1-154):** `tblHeader` em {xml.count('<w:tblHeader/>')} tabela(s); `cantSplit` em {xml.count('<w:cantSplit/>')} linhas.",
      "- **Saída da v1, na lixeira:**"] + [f"  - `{k}`: `{v}`;" for k, v in ANTERIOR.items()] + [
      f"- **Frases prescritas do suplemento conferidas intactas:** {', '.join(PRESCRITAS)}.",
      f"- **Seções:** {', '.join(orient)} (tabelas largas deitadas). **Número de página em todas as seções:** {'sim' if rodape_ok else 'não'}.",
      f"- **Conferências:** bordas verticais ativas {len(ruins)}; sombreamento {len(sombra)}; texto justificado {'sim' if justif else 'não'}; ordem dos títulos das tabelas {tit_tab}; remissões no documento {citadas}; identificação, marcações e caminhos {'ausentes' if not vaz else 'PRESENTES'}.",
      "", "## Saídas (fora do git)", "", f"- `{OUT}`: `{sha(OUT)}`;", "", f"- **Resultado:** {'FALHA: ' + '; '.join(falhas) if falhas else 'ok'}.", ""]
open(arq, "w", encoding="utf-8").write("\n".join(md))
print("gravado:", arq, "| falhas:", falhas, "| refs:", len(ordem), "| parágrafos:", n_par, "| seções:", orient)
if falhas:
    sys.exit(1)
